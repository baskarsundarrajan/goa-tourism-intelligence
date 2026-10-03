"""
run.py  —  CLI entry point for the Goa Tourism Intelligence system.

Keyless deterministic self-check (NO model, NO API key) — proves the whole pipeline
(DB -> analyzers -> rules engine -> manpower/resource plan -> .docx brief):

    python run.py --selfcheck

Full multi-agent run for one task (needs a model backend configured):

    python run.py --task events --month 12
    python run.py --task sustainability --site "Morjim Beach" --month 12
    python run.py --task business --district "South Goa" --category eco_stay
    python run.py --task itinerary --days 4 --zone south --pace balanced
    python run.py --task seasonal --month 12 --zone both --backend ollama

Outputs land in output/ (a combined .docx brief) and findings/ (per-task agent notes).
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")  # emoji/em-dash safe on Windows consoles
except (AttributeError, ValueError):
    pass

PROJECT_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = PROJECT_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)


def _load_dotenv() -> None:
    env = PROJECT_DIR / ".env"
    if not env.exists():
        return
    for raw in env.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        k, v = k.strip(), v.strip().strip('"').strip("'")
        if k and k not in os.environ:
            os.environ[k] = v


# ---------------------------------------------------------------------------
# Keyless deterministic self-check
# ---------------------------------------------------------------------------
def selfcheck() -> int:
    import domains
    import briefing
    from data import db

    print("=== GOA TOURISM — DETERMINISTIC SELF-CHECK (no model / no API key) ===\n")
    print(db.build())

    # default scenario per task -> the combined brief
    base = {
        "seasonal": {"month": 12, "zone": "both"},
        "itinerary": {"days": 4, "zone": "both", "interests": ["beach", "heritage", "nature"],
                      "month": 12, "pace": "balanced"},
        "business": {"district": "North Goa", "category": "beach_shack"},
        "sustainability": {"site": "Morjim Beach", "month": 12},
        "events": {"month": 12},
    }
    results = {}
    print("\n--- task verdicts (December peak scenario) ---")
    for k, p in base.items():
        r = domains.analyze(k, p)
        results[k] = r
        gates = ",".join(g["key"] for g in r["gates_fired"]) or "-"
        print(f"  {k:14s} idx={r['index']:5.1f}  {r['band']['key']:7s}  gates={gates}")

    # scenario-sensitivity checks (the recommendation must move with the scenario)
    print("\n--- scenario-sensitivity checks ---")
    jul_ev = domains.analyze("events", {"month": 7})["index"]
    dec_ev = results["events"]["index"]
    print(f"  events: July={jul_ev} vs December={dec_ev}")
    assert dec_ev < jul_ev, "events index should be worse in December than July"

    jul_su = domains.analyze("sustainability", {"site": "Morjim Beach", "month": 7})["index"]
    dec_su = results["sustainability"]["index"]
    print(f"  sustainability (Morjim): July={jul_su} vs December={dec_su}")
    assert dec_su < jul_su, "Morjim should be worse in December (over capacity) than July"

    colva = domains.analyze("sustainability", {"site": "Colva Beach", "month": 10})["index"]
    morjim = domains.analyze("sustainability", {"site": "Morjim Beach", "month": 10})["index"]
    print(f"  sensitivity (Oct): Colva(medium)={colva} vs Morjim(high)={morjim}")
    assert morjim < colva, "high-sensitivity site should score lower than medium at equal month"

    # build the combined .docx brief
    out = OUTPUT_DIR / "goa_tourism_policy_brief.docx"
    briefing.write_docx(str(out), results,
                        {"title": "Goa Tourism Intelligence — December Peak Readiness"})
    from docx import Document
    doc = Document(str(out))
    print(f"\n.docx written: {out} ({out.stat().st_size:,} bytes; "
          f"{len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables)")
    print("\nSELF-CHECK PASSED ✅  (deterministic core + brief verified, keyless)")
    return 0


# ---------------------------------------------------------------------------
# Full multi-agent run
# ---------------------------------------------------------------------------
def run_task(domain_key: str, params: dict, recursion_limit: int) -> int:
    import agent
    import briefing
    from model_backend import describe

    print(f"Model backend: {describe()}")
    print(f"Running '{domain_key}' with params {params} …\n")
    res = agent.run_domain(domain_key, params, recursion_limit=recursion_limit)

    print(f"VERDICT : {res['band']['label']}  (index {res['index']}/100)")
    print(f"HEADLINE: {res['headline']}")
    if res.get("gates_fired"):
        for g in res["gates_fired"]:
            print(f"  GATE: {g['key']} caps at {g['ceiling']}")
    total = sum(int(x["recommended"]) for x in res["manpower"])
    print(f"MANPOWER: ~{total:,} personnel across {len(res['manpower'])} roles")
    if res.get("agent_summary"):
        print("\nAgent summary:\n  " + res["agent_summary"])
    if res.get("_agent_error"):
        print(f"\n[agent fell back to deterministic result: {res['_agent_error']}]")

    out = OUTPUT_DIR / f"goa_{domain_key}_brief.docx"
    briefing.write_docx(str(out), {domain_key: res},
                        {"title": f"Goa Tourism — {res['title']}"})
    print(f"\n.docx brief: {out}")
    return 0


def _params_from_args(domain_key: str, a) -> dict:
    if domain_key == "seasonal":
        return {"month": a.month or 12, "zone": a.zone or "both"}
    if domain_key == "itinerary":
        interests = ([x.strip() for x in a.interests.split(",")] if a.interests
                     else ["beach", "heritage", "nature"])
        return {"days": a.days or 3, "zone": a.zone or "both", "interests": interests,
                "month": a.month or 12, "pace": a.pace or "balanced"}
    if domain_key == "business":
        return {"district": a.district or "North Goa", "category": a.category or "homestay"}
    if domain_key == "sustainability":
        return {"site": a.site or "Calangute Beach", "month": a.month or 12}
    if domain_key == "events":
        return {"month": a.month or 12}
    return {}


def main() -> int:
    ap = argparse.ArgumentParser(description="Goa Tourism Intelligence — multi-agent decision support.")
    ap.add_argument("--selfcheck", action="store_true", help="Keyless deterministic check + brief.")
    ap.add_argument("--task", choices=["seasonal", "itinerary", "business", "sustainability", "events"])
    ap.add_argument("--month", type=int)
    ap.add_argument("--zone")
    ap.add_argument("--site")
    ap.add_argument("--district")
    ap.add_argument("--category")
    ap.add_argument("--days", type=int)
    ap.add_argument("--pace")
    ap.add_argument("--interests", help="comma-separated, e.g. beach,heritage,wellness")
    ap.add_argument("--backend", choices=["anthropic", "ollama", "gemini"])
    ap.add_argument("--model")
    ap.add_argument("--recursion-limit", type=int, default=60)
    a = ap.parse_args()

    _load_dotenv()
    if a.backend:
        os.environ["GOA_MODEL_BACKEND"] = a.backend
    if a.model:
        os.environ["GOA_MODEL_ID"] = a.model

    if a.selfcheck:
        return selfcheck()
    if not a.task:
        ap.error("Provide --task {seasonal|itinerary|business|sustainability|events} or --selfcheck.")
    try:
        return run_task(a.task, _params_from_args(a.task, a), a.recursion_limit)
    except Exception as e:  # noqa: BLE001
        print(f"\n[ERROR] {type(e).__name__}: {e}", file=sys.stderr)
        print("If backend=anthropic, set ANTHROPIC_API_KEY. If backend=ollama, run Ollama and "
              "`ollama pull qwen2.5:7b`. Or use --selfcheck for the keyless path.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
