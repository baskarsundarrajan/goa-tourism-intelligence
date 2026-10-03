"""
briefing.py  —  compiles the task results into a Government-of-Goa policy brief (.docx).

One brief can cover any subset of the five tasks the user has run. For each task it
renders: a verdict banner, the core-rule scorecard, the sized manpower & resource tables,
the recommended actions and top risks. It closes with an aggregate manpower roll-up across
all included tasks, and a sources/disclaimer footer.
"""

from __future__ import annotations

import datetime as _dt

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from data import seed_data as S

_BAND_FILL = {"green": ("1B7F3B", "FFFFFF"), "amber": ("B8860B", "FFFFFF"),
              "orange": ("C05A1E", "FFFFFF"), "red": ("B22222", "FFFFFF")}
_RISK_INK = {"H": "B22222", "M": "B8860B", "L": "1B7F3B"}


def _shade(el, hex_fill: str) -> None:
    pr = el.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def _cell_shade(cell, hex_fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _run(p, text, *, size=9, bold=False, italic=False, color=None):
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _tight(p, before=0, after=2):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.0
    return p


def _trim(text: str, n: int) -> str:
    text = " ".join((text or "").split())
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def _score_ink(v: float) -> str:
    return "1B7F3B" if v >= 60 else "B8860B" if v >= 42 else "B22222"


def _table(doc, headers, rows, widths, head_fill="2A3542"):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl.style = "Table Grid"
    for cell, label in zip(tbl.rows[0].cells, headers):
        _cell_shade(cell, head_fill)
        _tight(cell.paragraphs[0], after=0)
        _run(cell.paragraphs[0], label, size=8, bold=True, color="FFFFFF")
    for row in rows:
        cells = tbl.add_row().cells
        for i, val in enumerate(row):
            _tight(cells[i].paragraphs[0], after=0)
            _run(cells[i].paragraphs[0], str(val), size=8)
    for row in tbl.rows:
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
    return tbl


def _task_section(doc, res: dict) -> None:
    band = res.get("band", {})
    bkey = band.get("key", "red")
    fill, ink = _BAND_FILL.get(bkey, ("444444", "FFFFFF"))

    p = _tight(doc.add_paragraph(), before=8, after=1)
    _run(p, f"{res.get('icon','')}  {res.get('title','')}", size=13, bold=True)
    p = _tight(doc.add_paragraph(), after=3)
    _run(p, f"Scenario: {res.get('scenario','')}", size=8.5, italic=True, color="666666")

    # verdict banner
    p = doc.add_paragraph(); _tight(p, before=1, after=2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _shade(p._p, fill)
    _run(p, f"{band.get('label','')}   ·   Preparedness index {res.get('index','-')}/100",
         size=10, bold=True, color=ink)

    p = _tight(doc.add_paragraph(), after=3)
    _run(p, band.get("headline", ""), size=9)

    # scorecard
    p = _tight(doc.add_paragraph(), before=3, after=1)
    _run(p, "Core-rule scorecard", size=9.5, bold=True)
    tbl = _table(doc, ("Dimension", "Wt", "Score", "Basis"),
                 [(d["name"], f"{d['weight']:.2f}", f"{d['score']:.0f}",
                   _trim(d.get("rationale", ""), 150)) for d in res.get("dimensions", [])],
                 (Inches(2.0), Inches(0.4), Inches(0.55), Inches(4.0)))
    # colour the score cells
    for i, d in enumerate(res.get("dimensions", []), start=1):
        cell = tbl.rows[i].cells[2]
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(_score_ink(d["score"]))
        cell.paragraphs[0].runs[0].bold = True

    # gates
    for g in res.get("gates_fired", []) or []:
        p = _tight(doc.add_paragraph(), before=2, after=1)
        _run(p, f"⚑ Gate fired — caps index at {g.get('ceiling')}: ", size=8.5, bold=True, color="B22222")
        _run(p, _trim(g.get("description", ""), 180), size=8.5)
        p = _tight(doc.add_paragraph(), after=1)
        _run(p, "Action: ", size=8.5, bold=True, color="444444")
        _run(p, _trim(g.get("action", ""), 200), size=8.5, italic=True)

    # manpower + resources
    mp = res.get("manpower", [])
    if mp:
        total = sum(int(x["recommended"]) for x in mp)
        p = _tight(doc.add_paragraph(), before=3, after=1)
        _run(p, f"Recommended manpower deployment  (total ≈ {total:,})", size=9.5, bold=True)
        _table(doc, ("Role", "Recommended", "Planning basis"),
               [(x.get("role", ""), x["recommended"], _trim(x.get("basis", ""), 90)) for x in mp],
               (Inches(2.6), Inches(1.0), Inches(3.4)))
    rs = res.get("resources", [])
    if rs:
        p = _tight(doc.add_paragraph(), before=3, after=1)
        _run(p, "Recommended resources / infrastructure", size=9.5, bold=True)
        _table(doc, ("Item", "Recommended", "Planning basis"),
               [(x.get("item", ""), x["recommended"], _trim(x.get("basis", ""), 90)) for x in rs],
               (Inches(2.6), Inches(1.0), Inches(3.4)))

    # recommendations
    recs = res.get("recommendations", [])
    if recs:
        p = _tight(doc.add_paragraph(), before=3, after=1)
        _run(p, "Policy actions", size=9.5, bold=True)
        for c in recs[:6]:
            b = _tight(doc.add_paragraph(style="List Bullet"), after=0)
            _run(b, _trim(c, 220), size=8.5)

    # risks
    risks = res.get("risks", [])
    if risks:
        p = _tight(doc.add_paragraph(), before=3, after=1)
        _run(p, "Key risks", size=9.5, bold=True)
        for r in risks[:4]:
            b = _tight(doc.add_paragraph(style="List Bullet"), after=0)
            _run(b, f"[{r.get('likelihood','-')}/{r.get('impact','-')}] ", size=8.5, bold=True,
                 color=_RISK_INK.get((r.get("impact") or "M").upper(), "B8860B"))
            _run(b, _trim(r.get("risk", ""), 110), size=8.5, bold=True)
            _run(b, " — " + _trim(r.get("mitigation", ""), 130), size=8.5, color="333333")

    # agent narrative (if the multi-agent path was used)
    if res.get("agent_summary"):
        p = _tight(doc.add_paragraph(), before=3, after=1)
        _run(p, "Multi-agent analyst summary", size=9, bold=True, color="444444")
        p = _tight(doc.add_paragraph(), after=2)
        _run(p, _trim(res["agent_summary"], 600), size=8.5, italic=True, color="333333")


def write_docx(path: str, results: dict[str, dict], meta: dict | None = None) -> None:
    meta = meta or {}
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.6)
    sec.left_margin = sec.right_margin = Inches(0.7)
    normal = doc.styles["Normal"]; normal.font.name = "Calibri"; normal.font.size = Pt(9)

    # title block
    p = _tight(doc.add_paragraph(), after=0)
    _run(p, "GOA TOURISM — POLICY & RESOURCE BRIEF", size=8, bold=True, color="666666")
    p = _tight(doc.add_paragraph(), after=1)
    _run(p, meta.get("title", "Goa Tourism Intelligence — Decision-Support Brief"), size=16, bold=True)
    p = _tight(doc.add_paragraph(), after=4)
    _run(p, f"Prepared for: {meta.get('prepared_for', 'Government of Goa — Department of Tourism')}"
            f"   ·   Generated: {_dt.datetime.now():%Y-%m-%d %H:%M}   ·   Evidence as of {S.AS_OF}",
         size=8, color="666666")

    # executive line
    ordered = [k for k in ("seasonal", "itinerary", "business", "sustainability", "events")
               if k in results]
    reds = [results[k]["title"] for k in ordered if results[k].get("band", {}).get("key") == "red"]
    p = _tight(doc.add_paragraph(), before=2, after=1)
    _run(p, "At a glance", size=10, bold=True)
    p = _tight(doc.add_paragraph(), after=3)
    _run(p, f"This brief covers {len(ordered)} task(s). "
            + (f"CRITICAL attention required: {', '.join(reds)}. " if reds
               else "No task is in the critical band. ")
            + "Each task below carries a transparent preparedness score and a sized manpower / "
              "resource plan the Government can act on.", size=9)

    for k in ordered:
        _task_section(doc, results[k])

    # aggregate manpower roll-up
    agg: dict[str, int] = {}
    for k in ordered:
        for x in results[k].get("manpower", []):
            agg[x.get("role", "")] = agg.get(x.get("role", ""), 0) + int(x["recommended"])
    if agg:
        p = _tight(doc.add_paragraph(), before=8, after=1)
        _run(p, "Aggregate manpower roll-up (across included tasks)", size=11, bold=True)
        total = sum(agg.values())
        _table(doc, ("Role", "Total recommended"),
               sorted(agg.items(), key=lambda kv: -kv[1]),
               (Inches(4.0), Inches(2.0)))
        p = _tight(doc.add_paragraph(), before=2, after=2)
        _run(p, f"Indicative total field deployment across these tasks ≈ {total:,} personnel. "
                f"Figures are planning estimates from open, editable norms — tune to Goa Police, "
                f"Drishti Marine and Directorate of Tourism doctrine before budgeting.", size=8.5,
             italic=True, color="444444")

    # footer
    p = _tight(doc.add_paragraph(), before=6, after=0)
    _run(p, "Sources incl.: ", size=7.5, bold=True, color="666666")
    _run(p, "; ".join(S.SOURCES[:4]) + "; …", size=7.5, color="666666")
    p = _tight(doc.add_paragraph(), after=0)
    _run(p, S.DISCLAIMER, size=7.5, italic=True, color="888888")
    p = _tight(doc.add_paragraph(), after=0)
    _run(p, "Decision-support generated by a multi-agent system (deterministic rules engine + "
            "LangChain deepagents). NOT an official Government of Goa position; the human "
            "decision-maker decides.", size=7.5, italic=True, color="888888")

    doc.save(path)
